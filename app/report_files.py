"""Build Excel (.xlsx) and PDF bytes for Admin reports."""
from __future__ import annotations

import io
from datetime import datetime

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _stamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def workbook_bytes(title: str, sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    first = True
    for sheet_name, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = (sheet_name or "Sheet")[:31]
        ws["A1"] = "PharmaLink"
        ws["A1"].font = Font(bold=True, size=14, color="1E3A8A")
        ws["A2"] = title
        ws["A3"] = "Generated " + _stamp()
        start = 5
        for col, h in enumerate(headers, 1):
            cell = ws.cell(start, col, h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left")
        for r_i, row in enumerate(rows, start + 1):
            for c_i, val in enumerate(row, 1):
                cell = ws.cell(r_i, c_i, val)
                cell.border = thin
        for col in range(1, max(len(headers), 1) + 1):
            letter = get_column_letter(col)
            longest = len(str(headers[col - 1])) if col <= len(headers) else 10
            for row in rows[:80]:
                if col <= len(row):
                    longest = max(longest, min(42, len(str(row[col - 1]))))
            ws.column_dimensions[letter].width = min(44, max(12, longest + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _ReportPdf(FPDF):
    def __init__(self, title: str):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.report_title = title
        self.set_auto_page_break(auto=True, margin=14)

    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.set_text_color(30, 58, 138)
        self.cell(0, 7, "PharmaLink", ln=1)
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(15, 23, 42)
        self.cell(0, 6, self.report_title, ln=1)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(100, 116, 139)
        self.cell(0, 5, "Generated " + _stamp(), ln=1)
        self.ln(2)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(148, 163, 184)
        self.cell(0, 8, f"Page {self.page_no()}", align="C")


def _ascii(val) -> str:
    text = "" if val is None else str(val)
    return text.encode("latin-1", "replace").decode("latin-1")


def pdf_bytes(title: str, headers: list[str], rows: list[list], intro: str = "") -> bytes:
    pdf = _ReportPdf(title)
    pdf.add_page()
    if intro:
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(71, 85, 105)
        pdf.multi_cell(0, 5, _ascii(intro))
        pdf.ln(2)
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    n = max(len(headers), 1)
    widths = [usable / n] * n
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_fill_color(30, 58, 138)
    pdf.set_text_color(255, 255, 255)
    for i, h in enumerate(headers):
        pdf.cell(widths[i], 7, _ascii(h)[:40], border=0, fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 7.5)
    pdf.set_text_color(30, 41, 59)
    fill = False
    for row in rows:
        if pdf.get_y() > pdf.h - 20:
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 8)
            pdf.set_fill_color(30, 58, 138)
            pdf.set_text_color(255, 255, 255)
            for i, h in enumerate(headers):
                pdf.cell(widths[i], 7, _ascii(h)[:40], border=0, fill=True)
            pdf.ln()
            pdf.set_font("Helvetica", "", 7.5)
            pdf.set_text_color(30, 41, 59)
        if fill:
            pdf.set_fill_color(241, 245, 249)
        else:
            pdf.set_fill_color(255, 255, 255)
        for i in range(n):
            val = row[i] if i < len(row) else ""
            pdf.cell(widths[i], 6, _ascii(val)[:48], border=0, fill=True)
        pdf.ln()
        fill = not fill
    out = pdf.output()
    return bytes(out) if isinstance(out, (bytes, bytearray)) else out.encode("latin-1")
